import json
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]


def output_json(result) -> dict:
    assert result.stdout, result.stderr
    return json.loads(result.stdout)


def test_model_export_writes_xlsx_with_both_sheets(tmp_path, run_cli):
    from openpyxl import load_workbook

    assert run_cli("init", str(tmp_path)).returncode == 0
    config_path = ROOT / "examples/ridgeline/config.yaml"
    out_path = tmp_path / "model.xlsx"

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(config_path),
        "--out",
        str(out_path),
    )

    assert result.returncode == 0
    payload = output_json(result)
    assert payload["ok"] is True
    assert payload["command"] == "model-export"
    data = payload["data"]
    assert set(data["sheets"]) == {"Assumptions", "Model"}
    assert out_path.exists()
    wb = load_workbook(out_path)
    assert set(wb.sheetnames) == {"Assumptions", "Model"}


def test_model_export_fails_without_workspace(tmp_path, run_cli):
    config_path = ROOT / "examples/ridgeline/config.yaml"
    out_path = tmp_path / "model.xlsx"

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(config_path),
        "--out",
        str(out_path),
    )

    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False
    assert payload["error"]["type"] == "workspace_not_initialized"


def test_model_export_fails_on_invalid_config(tmp_path, run_cli):
    assert run_cli("init", str(tmp_path)).returncode == 0
    bad_config = tmp_path / "bad.yaml"
    bad_config.write_text("not: valid: entity: config\n")
    out_path = tmp_path / "model.xlsx"

    result = run_cli(
        "model-export",
        str(tmp_path),
        "--config",
        str(bad_config),
        "--out",
        str(out_path),
    )

    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False


@pytest.mark.parametrize("alias_kind", ["same", "relative", "source symlink", "output symlink", "hard link"])
def test_model_export_preserves_configuration_when_output_aliases_it(tmp_path, run_cli, alias_kind):
    assert run_cli("init", str(tmp_path)).returncode == 0
    config = tmp_path / "config.yaml"
    original = (ROOT / "examples/ridgeline/config.yaml").read_bytes()
    config.write_bytes(original)
    source, output = config, config
    try:
        if alias_kind == "relative":
            (tmp_path / "child").mkdir()
            output = tmp_path / "child" / ".." / config.name
        elif alias_kind == "source symlink":
            source = tmp_path / "config-link.yaml"
            source.symlink_to(config)
        elif alias_kind == "output symlink":
            output = tmp_path / "model.xlsx"
            output.symlink_to(config)
        elif alias_kind == "hard link":
            output = tmp_path / "model.xlsx"
            output.hardlink_to(config)
    except OSError as error:
        pytest.skip(f"Cannot create a local file link: {error}")

    result = run_cli("model-export", str(tmp_path), "--config", str(source), "--out", str(output))

    assert config.read_bytes() == original
    assert source.read_bytes() == original
    assert result.returncode == 1
    payload = output_json(result)
    assert payload["ok"] is False
    assert payload["error"]["type"] == "export_failed"
